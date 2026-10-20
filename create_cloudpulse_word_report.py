from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path.cwd()
OUT = ROOT / "outputs"
CHARTS = OUT / "visualizations"

# Save as a new report so the original stays untouched.
REPORT = OUT / "CloudPulse_Model_Evaluation_Report_v2.docx"

if REPORT.exists():
    raise FileExistsError(
        f"{REPORT} already exists. Rename it or change REPORT in this script."
    )

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
for style_name in ("Title", "Heading 1", "Heading 2"):
    styles[style_name].font.name = "Aptos Display"

styles["Title"].font.size = Pt(30)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(31, 78, 121)
styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
styles["Heading 1"].font.size = Pt(18)
styles["Heading 2"].font.color.rgb = RGBColor(46, 116, 181)

def shade_cell(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_text(cell, value, bold=False):
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.bold = bold

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, heading in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], heading, True)
        shade_cell(table.rows[0].cells[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph()
    return table

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

def add_chart(filename, caption):
    path = CHARTS / filename
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(6.2))
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cp.runs:
            run.italic = True
            run.font.size = Pt(8)
    else:
        doc.add_paragraph(f"Chart not found: {path}")

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\n")
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("CloudPulse")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Cloud Resource Anomaly Detection\nModel Evaluation Report")
r.bold = True
r.font.size = Pt(17)
r.font.color.rgb = RGBColor(70, 70, 70)

doc.add_paragraph("\n")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Comparative evaluation of machine-learning and time-series approaches")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Dataset: 10,000 synthetic cloud-metric observations\n")
p.add_run("Evaluation focus: detection quality, false alarms, and robustness")
doc.add_page_break()

# Executive summary
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "CloudPulse is a learning and experimentation project for detecting unusual "
    "patterns in cloud-resource metrics. It uses a synthetic dataset containing "
    "CPU utilization, memory utilization, network throughput, timestamps, and "
    "anomaly labels. This report compares an engineered-feature XGBoost model, "
    "a threshold-based hybrid, LSTM and GRU sequence models, an adaptive "
    "network detector, and an ARIMA-based detector."
)
doc.add_paragraph(
    "On the recorded comparison, the CloudPulse V3 Hybrid achieved the highest "
    "F1-score (0.9614) and 93.33% recall with an observed false-alarm rate of "
    "approximately 0.07%. LSTM achieved 97.50% recall but produced more false "
    "alarms. Stress tests revealed that the hybrid's fixed network threshold "
    "can become unreliable when normal traffic patterns shift. The hybrid "
    "should therefore be treated as an experimental baseline, not a "
    "production-ready detector."
)

doc.add_heading("2. Project Objectives", level=1)
for item in [
    "Detect anomalous patterns in CPU, memory, and network metrics.",
    "Compare supervised feature-based detection with sequence and time-series approaches.",
    "Evaluate precision, recall, F1-score, and false-alarm rate.",
    "Examine detection by anomaly type and under altered traffic conditions.",
    "Document limitations before considering integration into a cloud-monitoring workflow.",
]:
    bullet(item)

doc.add_heading("3. Dataset and Preparation", level=1)
add_table(
    ["Property", "Description"],
    [
        ["Dataset size", "10,000 synthetic observations"],
        ["Observation interval", "5 minutes"],
        ["Metrics", "CPU %, memory %, network Mbps"],
        ["Labels", "is_anomaly and anomaly_type"],
        ["Normal observations", "9,750"],
        ["Anomalous observations", "250"],
        ["Anomaly categories", "memory drift (100), combined (60), network drop (50), CPU spike (40)"],
        ["Data checks", "No missing values or duplicate timestamps reported; timestamps ordered"],
    ],
)
doc.add_paragraph(
    "The dataset is synthetic and does not represent measurements collected "
    "from a live AWS account. Results demonstrate behavior on generated data "
    "and should not be interpreted as real-world cloud performance."
)

doc.add_heading("4. Models and Detection Strategies", level=1)
doc.add_heading("4.1 XGBoost V3", level=2)
doc.add_paragraph(
    "The XGBoost model uses engineered metric features, lagged observations, "
    "rolling statistics, and network-relative features. The recorded configuration "
    "uses 300 estimators, maximum depth 4, learning rate 0.05, subsampling 0.8, "
    "column sampling 0.8, and a chronological train/validation/test split. "
    "The selected probability threshold was 0.38."
)
doc.add_heading("4.2 CloudPulse V3 Hybrid", level=2)
doc.add_paragraph(
    "The hybrid flags an observation when the XGBoost probability meets its "
    "threshold or network throughput falls below the fixed threshold of "
    "22.3734 Mbps. This rule improved network-drop detection on the original "
    "synthetic test data, but stress testing showed that it can over-alert "
    "when normal network traffic becomes lower than expected."
)
doc.add_heading("4.3 LSTM and GRU", level=2)
doc.add_paragraph(
    "LSTM and GRU sequence models use sequence prediction or reconstruction "
    "errors as anomaly signals. Their thresholds were selected using the "
    "99th percentile of validation normal errors. In the recorded test, they "
    "detected most anomalies but generated more false alarms than XGBoost alone."
)
doc.add_heading("4.4 Adaptive Network Detector V1", level=2)
doc.add_paragraph(
    "This detector uses a rolling median and median absolute deviation (MAD) "
    "to estimate the recent network baseline and flag unusual deviations. "
    "Its recorded test recall was very low, indicating that the chosen rolling "
    "window and threshold strategy did not detect the injected anomalies well."
)
doc.add_heading("4.5 ARIMA Network Detector V1", level=2)
doc.add_paragraph(
    "The ARIMA(2,1,2) experiment uses sequential one-step-ahead network forecasts "
    "and flags large absolute forecast errors. It detected some combined anomalies "
    "but missed most CPU and memory anomalies because it primarily models network "
    "throughput. Its overall recall was low."
)

doc.add_heading("5. Comparative Test Results", level=1)
add_table(
    ["Model", "Rows", "Precision", "Recall", "F1-score", "False-alarm rate"],
    [
        ["XGBoost V3 Model Only", "1,499", "98.82%", "70.00%", "81.95%", "0.07%"],
        ["CloudPulse V3 Hybrid", "1,499", "99.12%", "93.33%", "96.14%", "0.07%"],
        ["LSTM", "1,499", "89.31%", "97.50%", "93.23%", "1.02%"],
        ["GRU", "1,499", "87.12%", "95.83%", "91.27%", "1.23%"],
        ["Adaptive Network Detector V1", "1,499", "13.79%", "3.33%", "5.37%", "1.81%"],
        ["ARIMA Network Detector V1", "1,500", "58.06%", "15.00%", "23.84%", "0.94%"],
    ],
)
doc.add_paragraph(
    "Precision is the proportion of predicted anomalies that were actual "
    "anomalies; recall is the proportion of actual anomalies detected; F1-score "
    "is the harmonic mean of precision and recall; false-alarm rate is false "
    "positives divided by actual normal observations."
)
doc.add_paragraph(
    "Comparison caveat: the ARIMA result uses 1,500 test rows, whereas the "
    "other rows in the unified comparison use 1,499. Thresholds and detection "
    "mechanisms also differ, so this table records the current experiments "
    "rather than a fully controlled benchmark. Re-run every model on identical "
    "timestamp rows and with a consistent threshold-selection protocol before "
    "making definitive model-ranking claims."
)
add_chart(
    "model_performance_comparison.png",
    "Figure 1. Model performance comparison generated by the project's visualization script."
)
add_chart(
    "confusion_matrices.png",
    "Figure 2. Confusion matrices for the evaluated approaches."
)

doc.add_heading("6. Hybrid Anomaly-Type Analysis", level=1)
add_table(
    ["Anomaly type", "Detected", "Total", "Recall"],
    [
        ["Combined", "18", "20", "90.0%"],
        ["CPU spike", "19", "20", "95.0%"],
        ["Memory drift", "45", "50", "90.0%"],
        ["Network drop", "30", "30", "100.0%"],
    ],
)
doc.add_paragraph(
    "The hybrid's strongest result was on network drops in the original test "
    "data. This is consistent with its explicit network-throughput rule. "
    "That advantage must be interpreted alongside the stress-test findings."
)
add_chart(
    "anomaly_type_comparison.png",
    "Figure 3. Detection performance by anomaly type."
)

doc.add_heading("7. Robustness and Stress Testing", level=1)
doc.add_paragraph(
    "Additional experiments changed normal traffic levels or introduced "
    "overlapping and milder network-drop patterns. These tests revealed that "
    "the fixed network threshold is sensitive to traffic distribution."
)
add_table(
    ["Scenario", "Detector", "FP", "FN", "Recall", "F1-score", "FAR"],
    [
        ["Low normal traffic", "XGBoost", "48", "33", "72.50%", "68.24%", "3.48%"],
        ["Low normal traffic", "Hybrid", "276", "5", "95.83%", "45.01%", "20.01%"],
        ["Mild network drops", "XGBoost", "1", "36", "70.00%", "81.95%", "0.07%"],
        ["Mild network drops", "Hybrid", "1", "22", "81.67%", "89.50%", "0.07%"],
        ["Overlapping traffic", "XGBoost", "44", "33", "72.50%", "69.32%", "3.19%"],
        ["Overlapping traffic", "Hybrid", "186", "17", "85.83%", "50.37%", "13.49%"],
    ],
)
doc.add_paragraph(
    "The gated hybrid produced the same predictions as XGBoost-only on the "
    "recorded baseline, low-normal-traffic, mild-network-drop, and overlapping-"
    "traffic scenarios. It therefore did not provide a measurable improvement "
    "in these experiments. A relative network rule also introduced additional "
    "false positives without increasing network-drop detections in its recorded test."
)
doc.add_paragraph(
    "Recommendation: do not deploy the fixed network threshold unchanged. "
    "Future experiments should evaluate thresholds using chronological "
    "validation data, assess normal-traffic drift, report per-scenario false "
    "alarms, and compare all models on identical timestamps."
)

doc.add_heading("8. Limitations", level=1)
for item in [
    "The dataset is synthetic and may not represent real cloud workloads.",
    "The hybrid relies on a fixed network threshold that is sensitive to traffic shifts.",
    "Model thresholds and detection methods are not fully identical across experiments.",
    "The ARIMA experiment focuses on network throughput and cannot independently capture every CPU or memory anomaly.",
    "Performance on unseen real-world workloads, different AWS regions, or production monitoring data has not been established.",
    "High test scores do not establish that the system is ready for production deployment.",
]:
    bullet(item)

doc.add_heading("9. Conclusion and Next Steps", level=1)
doc.add_paragraph(
    "The experiments show that different detectors provide different trade-offs. "
    "The hybrid achieved the strongest recorded F1-score on the original synthetic "
    "test set, while LSTM achieved higher recall than the hybrid at the cost of "
    "more false alarms. The adaptive detector and ARIMA experiment were weak "
    "overall in their current configurations. Stress testing exposed an important "
    "reliability issue with the hybrid's fixed network threshold."
)
for item in [
    "Standardize a common timestamp-based test set for every model.",
    "Use validation-only threshold selection and avoid tuning on the final test set.",
    "Repeat stress tests across several traffic levels and anomaly severities.",
    "Explore adaptive or context-aware network thresholds with explicit false-alarm constraints.",
    "Evaluate on safely collected, anonymized cloud telemetry before making real-world claims.",
    "Keep the hybrid as an experimental benchmark until its robustness improves.",
]:
    bullet(item)

doc.add_heading("10. Reproducibility and Project Artifacts", level=1)
doc.add_paragraph(
    "This report was generated from the CloudPulse project workspace. The "
    "visualizations are expected under outputs/visualizations/. The project also "
    "contains model-specific scripts, prediction files, metric outputs, and "
    "configuration files. Keep these artifacts together with the dataset and "
    "record dependency versions when preparing a final academic submission."
)

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CloudPulse | Model Evaluation Report")

doc.save(REPORT)
print(f"Report created successfully: {REPORT}")
print(f"File size: {REPORT.stat().st_size:,} bytes")
print("Charts folder:", CHARTS)




